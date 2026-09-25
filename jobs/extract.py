"""Extract text from brand docs, OCR scanned PDFs (BUILD_PLAN.md section 5).

Spec (worker/jobs/extract.ts buildSpec):
  docs: [{id, r2_key, mime_type, ext, text_key}]   files to read
  draft_profile: bool                                also draft the Brand Profile from all text
  done_text_keys: [r2 key]                           text of docs read earlier (for the draft)
  ocr_min_chars_per_page: int                        below this, a PDF page is treated as a picture
  llm: {key, model, system, sections} | null         OpenRouter call for the draft

Each doc's text is written to R2 at text_key. The result carries only a status, a character
count and a reason code per doc (never the text), plus the drafted profile sections when asked.
Logs: step names and counts only (section 13).
"""
from __future__ import annotations

import json
import logging
import os
import shutil
import subprocess
import tempfile
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

from common import WORK, Job, log, r2_client, r2_download, run

MIN_DOC_CHARS = 20
# pypdf warns on odd files; its messages stay out of the public log (section 13).
logging.getLogger("pypdf").setLevel(logging.ERROR)


class Unreadable(Exception):
    """A file we could not get words out of; `reason` is a code the Worker turns into a sentence."""

    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


# ---------- per-format readers ----------

def _ocr_pdf(path: Path) -> str:
    """Render each page to an image (poppler's pdftoppm) and read it with tesseract."""
    if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
        log("extract.ocr.missing_tools")
        return ""
    out: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-r", "200", "-png", str(path), f"{tmp}/p"], check=True, capture_output=True, timeout=900)
        pages = sorted(Path(tmp).glob("p*.png"))
        for i, img in enumerate(pages):
            res = subprocess.run(["tesseract", str(img), "stdout", "-l", "eng"], capture_output=True, timeout=300)
            out.append(res.stdout.decode("utf-8", errors="replace"))
        log("extract.ocr.pages", pages=len(pages))
    return "\n\n".join(out)


def read_pdf(path: Path, min_chars_per_page: int) -> tuple[str, bool]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except PdfReadError as e:
        raise Unreadable("corrupt") from e
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise Unreadable("locked")
        except Exception as e:  # noqa: BLE001
            raise Unreadable("locked") from e
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = "\n\n".join(pages)
    per_page = len(text.strip()) / max(1, len(pages))
    if per_page >= min_chars_per_page:
        return text, False
    # Scanned PDF: pages are pictures of text.
    ocr = _ocr_pdf(path)
    return (ocr if len(ocr.strip()) > len(text.strip()) else text), True


def read_docx(path: Path) -> str:
    import docx  # python-docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        d = docx.Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as e:
        raise Unreadable("corrupt") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_plain(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise Unreadable("corrupt")


def extract_text(path: Path, mime: str, ext: str, min_chars_per_page: int = 50) -> tuple[str, bool]:
    """Returns (text, used_ocr). Raises Unreadable with a reason code."""
    if mime == "application/pdf" or ext == "pdf":
        text, ocr = read_pdf(path, min_chars_per_page)
    elif ext == "docx" or mime.endswith("wordprocessingml.document"):
        text, ocr = read_docx(path), False
    elif mime.startswith("text/") or ext in ("md", "txt", "markdown"):
        text, ocr = read_plain(path), False
    else:
        raise Unreadable("unsupported")
    text = text.replace("\x00", "").strip()
    if len(text) < MIN_DOC_CHARS:
        raise Unreadable("no_text")
    return text, ocr


# ---------- OpenRouter (profile draft) ----------

def llm_json(key: str, model: str, system: str, user: str, max_tokens: int = 2500) -> dict[str, Any]:
    body = json.dumps({
        "model": model,
        "max_tokens": max_tokens,
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
        "response_format": {"type": "json_object"},
    }).encode()
    req = urllib.request.Request("https://openrouter.ai/api/v1/chat/completions", data=body, method="POST", headers={
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "X-Title": "Sheila Studio",
    })
    for attempt in range(5):
        try:
            with urllib.request.urlopen(req, timeout=180) as res:
                data = json.loads(res.read().decode())
            content = data["choices"][0]["message"]["content"] or ""
            start, end = content.find("{"), content.rfind("}")
            if start < 0 or end < 0:
                raise ValueError("no json in answer")
            return json.loads(content[start : end + 1])
        except urllib.error.HTTPError as e:
            log("llm.retry", status=e.code, attempt=attempt)
            if e.code in (401, 402, 403):
                raise
        except (urllib.error.URLError, TimeoutError, ValueError, KeyError, json.JSONDecodeError):
            log("llm.retry", attempt=attempt)
        time.sleep(min(60, 5 * 2 ** attempt))
    raise RuntimeError("llm unavailable")


CHUNK = 40_000


def draft_profile(llm: dict[str, Any], texts: list[str]) -> dict[str, str]:
    key = llm.get("key") or os.environ.get("OPENROUTER_API_KEY")
    if not key:
        raise RuntimeError("openrouter key missing")
    joined = [f"--- Document {i + 1} ---\n{t}" for i, t in enumerate(texts)]
    everything = "\n\n".join(joined)
    if len(everything) <= CHUNK:
        return llm_json(key, llm["model"], llm["system"], everything)
    # Too much for one call: condense each chunk into notes, then draft from the notes.
    notes: list[str] = []
    chunks = [everything[i : i + CHUNK] for i in range(0, len(everything), CHUNK)]
    for i, chunk in enumerate(chunks):
        part = llm_json(key, llm["model"], llm["system"], chunk)
        notes.append(json.dumps(part))
        log("extract.profile.chunk", done=i + 1, total=len(chunks))
    return llm_json(key, llm["model"], llm["system"], "Merge these partial profiles into one, keeping every concrete detail:\n\n" + "\n\n".join(notes))


# ---------- main ----------

def main(job: Job, spec: dict[str, Any]) -> dict[str, Any]:
    docs = spec.get("docs") or []
    min_cpp = int(spec.get("ocr_min_chars_per_page") or 50)
    work = WORK / "extract"
    work.mkdir(parents=True, exist_ok=True)
    results: list[dict[str, Any]] = []
    texts: list[str] = []
    bucket = os.environ.get("R2_BUCKET", "")
    s3 = r2_client() if docs or spec.get("done_text_keys") else None

    for i, d in enumerate(docs):
        job.progress("reading", i, len(docs))
        local = work / d["id"]
        try:
            r2_download(d["r2_key"], local)
        except Exception:  # noqa: BLE001
            results.append({"id": d["id"], "status": "unreadable", "char_count": 0, "reason": "missing"})
            log("extract.doc", ok=False, reason="missing")
            continue
        try:
            text, ocr = extract_text(local, d.get("mime_type", ""), d.get("ext", ""), min_cpp)
            s3.put_object(Bucket=bucket, Key=d["text_key"], Body=text.encode("utf-8"), ContentType="text/plain; charset=utf-8")
            results.append({"id": d["id"], "status": "done", "char_count": len(text), "ocr": ocr})
            texts.append(text)
            log("extract.doc", ok=True, chars=len(text), ocr=ocr)
        except Unreadable as e:
            results.append({"id": d["id"], "status": "unreadable", "char_count": 0, "reason": e.reason})
            log("extract.doc", ok=False, reason=e.reason)
        except Exception as e:  # noqa: BLE001
            results.append({"id": d["id"], "status": "unreadable", "char_count": 0, "reason": "corrupt"})
            log("extract.doc", ok=False, reason="error", kind=type(e).__name__)
        finally:
            local.unlink(missing_ok=True)
    job.progress("reading", len(docs), len(docs))

    out: dict[str, Any] = {"docs": results}
    if spec.get("draft_profile") and spec.get("llm"):
        for key in spec.get("done_text_keys") or []:
            try:
                obj = s3.get_object(Bucket=bucket, Key=key)
                texts.append(obj["Body"].read().decode("utf-8", errors="replace"))
            except Exception:  # noqa: BLE001
                log("extract.profile.text_missing")
        if texts:
            job.progress("drafting", 0, 1)
            profile = draft_profile(spec["llm"], texts)
            out["profile"] = {k: profile.get(k, "") for k in spec["llm"].get("sections", [])}
            log("extract.profile", sections=len(out["profile"]))
    log("extract.done", docs=len(results), read=sum(1 for r in results if r["status"] == "done"))
    return out


if __name__ == "__main__":
    run(main)
